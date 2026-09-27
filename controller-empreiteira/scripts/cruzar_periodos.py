#!/usr/bin/env python3
"""
Cruza totais por Unidade de Negocio x Mes/Ano entre duas bases (ex: Planilhao
bruto x Quadro Estrategico consolidado, ou duas exportacoes de periodos
diferentes) e aponta divergencias acima de uma tolerancia.

Funciona com dois formatos de entrada:
  (a) base "linha a linha" com colunas Unidade de Negocio / Ano / Mes / Valor
      (detectadas automaticamente, como no Planilhao);
  (b) base "pivotada" com uma coluna de rotulo de unidade de negocio e colunas
      de mes (JAN, FEV, MAR, ...) com o total do mes, como no Quadro
      Estrategico. Nesse caso passe --pivotado nesse arquivo.

Uso:
    python cruzar_periodos.py base1.xlsx base2.xlsx [--pivotado1] [--pivotado2] [--tolerancia 1.0]

O script NAO decide qual base esta certa -- ele so aponta onde os totais nao
batem para o controller investigar a causa (rateio nao refletido, lancamento
fora do periodo, erro de formula no consolidado, etc).
"""
import argparse
import json
import re
import unicodedata
import sys
from collections import defaultdict

from aw_comum import classificar_unidade

try:
    import openpyxl
except ImportError:
    print("Instale openpyxl: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(1)

MESES = {
    "JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12,
}


def _norm(t):
    t = unicodedata.normalize("NFKD", str(t)).encode("ascii", "ignore").decode()
    return t.lower().strip()


def col_index(header, *candidatos):
    lower = [_norm(h) if h is not None else "" for h in header]
    # 1) match exato primeiro (evita "ano" casar com "plano de contas")
    for cand in candidatos:
        for i, h in enumerate(lower):
            if h == _norm(cand):
                return i
    # 2) match por palavra inteira dentro do texto do cabecalho
    for cand in candidatos:
        padrao = re.compile(r"\b" + re.escape(_norm(cand)) + r"\b")
        for i, h in enumerate(lower):
            if padrao.search(h):
                return i
    return None


def unidade_chave(nome):
    """Nome canonico da unidade AW (grafia diferente nao vira unidade a mais)."""
    canon, _ = classificar_unidade(nome)
    return canon


def _mes_int(v):
    if isinstance(v, (int, float)):
        return int(v)
    t = _norm(v)[:3].upper()
    if t in MESES:
        return MESES[t]
    return int(t) if t.isdigit() else None


def _ano_int(v):
    try:
        return int(float(str(v).strip()))
    except ValueError:
        return None


def totais_linha_a_linha(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx_unidade = col_index(header, "unidade de negocio", "unidade")
    idx_mes = col_index(header, "mes")
    idx_ano = col_index(header, "ano")
    idx_valor = col_index(header, "valor")
    if None in (idx_unidade, idx_mes, idx_ano, idx_valor):
        raise ValueError(
            f"Nao encontrei colunas de Unidade/Mes/Ano/Valor em {path}. "
            f"Cabecalho lido: {header}. Se este arquivo for pivotado (tipo Quadro "
            f"Estrategico), rode com --pivotado."
        )
    totais = defaultdict(float)
    ignoradas = 0
    for r in rows:
        if not r:
            continue
        unidade, mes, ano, valor = r[idx_unidade], r[idx_mes], r[idx_ano], r[idx_valor]
        if unidade is None or mes is None or ano is None or not isinstance(valor, (int, float)):
            continue
        m, a = _mes_int(mes), _ano_int(ano)
        if m is None or a is None:
            ignoradas += 1
            continue
        totais[(unidade_chave(unidade), a, m)] += valor
    if ignoradas:
        print(f"AVISO: {ignoradas} linha(s) de {path} com Mes/Ano ilegivel foram ignoradas.", file=sys.stderr)
    return totais


def totais_pivotado(path, sheet=None, ano_forcado=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    # acha a linha de cabecalho com os meses (JAN, FEV, ...)
    header_row_idx = None
    for i, r in enumerate(rows):
        if r is None:
            continue
        textos = [str(c).strip().upper()[:3] for c in r if c is not None]
        if sum(1 for t in textos if t in MESES) >= 3:
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError(f"Nao encontrei uma linha de cabecalho com nomes de mes em {path}")

    header = rows[header_row_idx]
    mes_por_col = {}
    for j, c in enumerate(header):
        if c is None:
            continue
        chave = str(c).strip().upper()[:3]
        if chave in MESES:
            mes_por_col[j] = MESES[chave]

    # coluna de rotulo = primeira coluna de texto nao vazia antes das colunas de mes
    label_col = None
    for j in range(min(mes_por_col.keys())):
        for r in rows[header_row_idx + 1:header_row_idx + 5]:
            if r and j < len(r) and isinstance(r[j], str) and r[j].strip():
                label_col = j
                break
        if label_col is not None:
            break
    if label_col is None:
        label_col = 1  # fallback comum neste tipo de planilha

    totais = defaultdict(float)
    if ano_forcado is not None:
        ano_assumido = ano_forcado
    else:
        ano_assumido = None
        for r in rows:
            if r is None:
                continue
            m = re.search(r"(20\d{2})", " ".join(str(c) for c in r if c is not None))
            if m:
                ano_assumido = int(m.group(1))
                break
        # O ano e adivinhado pelo primeiro numero de 4 digitos tipo "20XX" encontrado na planilha
        # (ex: no titulo "Quadro Estrategico 2026") -- isso e uma heuristica fragil: uma celula com
        # outro numero de 4 digitos antes do titulo real, ou uma planilha com mais de um ano
        # misturado, engana essa deteccao. Por isso sempre avisamos o ano usado, e o script aceita
        # --ano1/--ano2 para o controller informar o ano explicitamente em vez de confiar no
        # palpite quando a base for ambigua.
        if ano_assumido is None:
            print(
                f"AVISO: nao consegui identificar o ano em {path} -- os totais desse arquivo vao "
                f"ficar com ano=None e provavelmente NAO vao casar com o outro arquivo. Rode de "
                f"novo passando --ano1/--ano2 com o ano correto.",
                file=sys.stderr,
            )
        else:
            print(
                f"Ano assumido para {path}: {ano_assumido} (detectado automaticamente -- "
                f"confirme com --ano1/--ano2 se estiver errado).",
                file=sys.stderr,
            )

    for r in rows[header_row_idx + 1:]:
        if r is None or label_col >= len(r):
            continue
        label = r[label_col]
        if not isinstance(label, str) or not label.strip():
            continue
        if "total" in _norm(label):
            continue  # linha TOTAL/subtotal do quadro: somar dobraria o valor
        for j, mes in mes_por_col.items():
            if j >= len(r):
                continue
            valor = r[j]
            if isinstance(valor, (int, float)):
                totais[(unidade_chave(label), ano_assumido, mes)] += valor

    return totais


def carregar_totais(path, pivotado, sheet=None, ano_forcado=None):
    if pivotado:
        return totais_pivotado(path, sheet, ano_forcado=ano_forcado)
    return totais_linha_a_linha(path, sheet)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("arquivo1")
    ap.add_argument("arquivo2")
    ap.add_argument("--sheet1", default=None)
    ap.add_argument("--sheet2", default=None)
    ap.add_argument("--pivotado1", action="store_true")
    ap.add_argument("--pivotado2", action="store_true")
    ap.add_argument("--ano1", type=int, default=None,
                    help="forca o ano do arquivo1 quando --pivotado1 (evita depender do palpite "
                         "automatico de ano)")
    ap.add_argument("--ano2", type=int, default=None,
                    help="forca o ano do arquivo2 quando --pivotado2")
    ap.add_argument("--tolerancia", type=float, default=1.0, help="diferenca absoluta tolerada")
    args = ap.parse_args()

    t1 = carregar_totais(args.arquivo1, args.pivotado1, args.sheet1, args.ano1)
    t2 = carregar_totais(args.arquivo2, args.pivotado2, args.sheet2, args.ano2)

    chaves = set(t1) | set(t2)
    divergencias = []
    for chave in sorted(chaves):
        v1 = t1.get(chave, 0.0)
        v2 = t2.get(chave, 0.0)
        diff = v1 - v2
        if abs(diff) > args.tolerancia:
            unidade, ano, mes = chave
            divergencias.append({
                "unidade": unidade, "ano": ano, "mes": mes,
                "valor_arquivo1": round(v1, 2), "valor_arquivo2": round(v2, 2),
                "diferenca": round(diff, 2),
            })

    resultado = {
        "arquivo1": args.arquivo1,
        "arquivo2": args.arquivo2,
        "tolerancia": args.tolerancia,
        "combinacoes_unidade_mes_arquivo1": len(t1),
        "combinacoes_unidade_mes_arquivo2": len(t2),
        "divergencias_encontradas": len(divergencias),
        "divergencias": divergencias,
    }
    print(json.dumps(resultado, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
