#!/usr/bin/env python3
"""
Painel do controller AW: projeto -> unidade de negocio -> consolidado.

Le o arquivo unico da AW (abas Faturamento, Projetos, CockpitRelatorioEmpreiteira e, se houver,
Planilhao) e monta, com rotulo de cada numero:
  - por projeto: faturado, recebido, retencoes, a receber, aging, custo incorrido (Cockpit e
    Planilhao), margem real e alertas;
  - por unidade: as mesmas somas + caixa real (conta bancaria) x caixa virtual (conta gerencial);
  - consolidado AW (13 ativas) com desativadas, controle a parte e fora do cadastro em linhas
    separadas;
  - controles: faturamento orfao, codigo duplicado no cadastro, unidade nao reconhecida, virtual que
    nao zera no consolidado, Cockpit x Planilhao, recebido > faturado, mobilizacao positiva, obra
    com margem negativa.

Nao altera o arquivo de entrada. Saida: resumo no terminal, JSON (--json) e planilha de trabalho
separada (--xlsx).

Uso:
    python painel_aw.py <arquivo.xlsx> [--planilhao outro.xlsx] [--data-base 30/09/2026]
                        [--aba-faturamento NOME] [--aba-projetos NOME] [--aba-cockpit NOME]
                        [--aba-planilhao NOME] [--tolerancia 1.0] [--json saida.json]
                        [--xlsx saida.xlsx]
"""
import argparse
import json
import sys
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    print("Instale openpyxl: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(1)

from aw_comum import (
    SITUACAO_ATIVA, SITUACAO_DESATIVADA, SITUACAO_A_PARTE, SITUACAO_FORA,
    achar_aba, classificar_unidade, col_index, eh_conta_bancaria, fmt_brl, get,
    letra_para_indice, ler_tabela, norm, parse_data, parse_numero, parse_projeto,
)

FAIXAS_AGING = ["A vencer", "1-30", "31-60", "61-90", ">90", "Sem vencimento"]


def faixa_aging(venc, data_base):
    if venc is None or data_base is None:
        return "Sem vencimento"
    d = (data_base - venc).days
    if d <= 0:
        return "A vencer"
    if d <= 30:
        return "1-30"
    if d <= 60:
        return "31-60"
    if d <= 90:
        return "61-90"
    return ">90"


def novo_projeto():
    return {
        "descritivo": None, "unidade": None, "situacao_unidade": None, "status": None,
        "no_cadastro": False,
        "faturado": 0.0, "recebido": 0.0, "retencoes": 0.0, "retencao_conhecida": False,
        "a_receber": 0.0, "aging": defaultdict(float), "n_nf": 0,
        "cockpit_mo": 0.0, "cockpit_materiais": 0.0, "cockpit_mobilizacao": 0.0,
        "cockpit_outros": 0.0, "no_cockpit": False,
        "planilhao_receita": 0.0, "planilhao_custo": 0.0, "no_planilhao": False,
        "clientes": defaultdict(float),
    }


# ------------------------------------------------------------------------ Projetos

def ler_projetos(ws, projetos, controles):
    header, corpo, _ = ler_tabela(ws)
    idx = col_index(header, "projeto", "projetos", "descricao")
    # A aba Projetos costuma ter uma coluna so, as vezes sem cabecalho: se o "cabecalho" ja parece
    # um projeto, ele tambem e dado.
    linhas = corpo
    if idx is None:
        idx = 0
        if parse_projeto(header[0]) and parse_projeto(header[0])["codigo"]:
            linhas = [tuple(header)] + list(corpo)
    vistos = defaultdict(set)
    for r in linhas:
        p = parse_projeto(get(r, idx))
        if not p or not p["codigo"]:
            if p:
                controles.append(ctrl("PENDÊNCIA", "Projetos", "registro sem código de projeto",
                                      p["descritivo"], None))
            continue
        unidade, situacao = classificar_unidade(p["unidade_texto"])
        vistos[p["codigo"]].add(unidade)
        pr = projetos.setdefault(p["codigo"], novo_projeto())
        pr.update({"descritivo": p["descritivo"], "unidade": unidade, "situacao_unidade": situacao,
                   "status": p["status"], "no_cadastro": True})
        if p["unidade_texto"] is None:
            controles.append(ctrl("PENDÊNCIA", "Projetos", "projeto sem '- AW <unidade>' no descritivo",
                                  p["codigo"], None))
        elif situacao == SITUACAO_FORA:
            controles.append(ctrl("INCONSISTÊNCIA", "Projetos", "unidade não reconhecida no cadastro AW "
                                  "(grafia ou unidade nova — confirmar)", f"{p['codigo']} → {unidade}", None))
    for cod, us in vistos.items():
        if len(us) > 1:
            controles.append(ctrl("INCONSISTÊNCIA", "Projetos", "mesmo código em unidades diferentes",
                                  f"{cod}: {sorted(us)}", None))
    base = {c.split("-")[0] for c in vistos if "-" in c}
    for b in sorted(base):
        if b in vistos:
            controles.append(ctrl("PENDÊNCIA", "Projetos", "código com e sem sufixo -NN (mesmo contrato "
                                  "ou aditivo? não unido)", b, None))
    return {"linhas_lidas": len(linhas), "codigos": len(vistos)}


# ------------------------------------------------------------------------ Faturamento

def ler_faturamento(ws, projetos, controles, data_base):
    header, corpo, _ = ler_tabela(ws)
    premissas = []
    i_proj = col_index(header, "projeto")
    i_bruto = col_index(header, "valor bruto", "faturado", "valor nf", "valor total")
    i_receb = col_index(header, "valor recebido liquido", "valor recebido", "recebido", "recebimento")
    i_liq = col_index(header, "valor liquido", evitar=("recebido",))
    i_ret = [i for i, h in enumerate(header) if "reten" in norm(h)]
    i_venc = col_index(header, "vencimento", "data vencimento", "dt vencimento")
    i_emis = col_index(header, "emissao", "data emissao", "dt emissao", "data nf")
    i_cli = col_index(header, "cliente", "tomador", "razao social")
    i_sit = col_index(header, "situacao", "status nf", "status")

    # fallback para o layout do modelo (W = Projeto, I = Valor Bruto, U = Valor Recebido Liquido)
    for nome, atual, letra in (("Projeto", i_proj, "W"), ("Valor Bruto", i_bruto, "I"),
                               ("Valor Recebido Líquido", i_receb, "U")):
        if atual is None:
            premissas.append(f"coluna '{nome}' não achada pelo rótulo — usada a coluna {letra} do "
                             f"modelo. CONFIRMAR.")
    if i_proj is None:
        i_proj = letra_para_indice("W")
    if i_bruto is None:
        i_bruto = letra_para_indice("I")
    if i_receb is None:
        i_receb = letra_para_indice("U")
    if i_liq is None and not i_ret:
        premissas.append("sem coluna de retenção/valor líquido: A receber = Bruto − Recebido Líquido "
                         "INCLUI retenções (superestima o CR). Rotular PREMISSA.")

    orfao = defaultdict(float)
    canceladas = 0
    texto_num = 0
    apos_corte = 0
    for r in corpo:
        if i_sit is not None and "cancel" in norm(get(r, i_sit)):
            canceladas += 1
            continue
        emis = parse_data(get(r, i_emis)) if i_emis is not None else None
        if data_base and emis and emis > data_base:
            apos_corte += 1
            continue
        bruto, t1 = parse_numero(get(r, i_bruto))
        receb, t2 = parse_numero(get(r, i_receb))
        texto_num += int(t1) + int(t2)
        bruto = bruto or 0.0
        receb = receb or 0.0
        ret = None
        if i_liq is not None:
            liq, _ = parse_numero(get(r, i_liq))
            if liq is not None:
                ret = bruto - liq
        elif i_ret:
            ret = sum((parse_numero(get(r, i))[0] or 0.0) for i in i_ret)
        p = parse_projeto(get(r, i_proj))
        cod = p["codigo"] if p and p["codigo"] else None
        if cod is None:
            orfao["(sem código de projeto)"] += bruto
            continue
        pr = projetos.setdefault(cod, novo_projeto())
        if not pr["no_cadastro"]:
            orfao[cod] += bruto
            if pr["descritivo"] is None and p:
                u, s = classificar_unidade(p["unidade_texto"])
                pr.update({"descritivo": p["descritivo"], "unidade": u, "situacao_unidade": s,
                           "status": p["status"]})
        pr["n_nf"] += 1
        pr["faturado"] += bruto
        pr["recebido"] += receb
        if ret is not None:
            pr["retencoes"] += abs(ret)
            pr["retencao_conhecida"] = True
        aberto = bruto - receb - (abs(ret) if ret is not None else 0.0)
        if aberto > 0.005:
            venc = parse_data(get(r, i_venc)) if i_venc is not None else None
            pr["aging"][faixa_aging(venc, data_base)] += aberto
        if i_cli is not None:
            pr["clientes"][str(get(r, i_cli) or "(sem cliente)").strip()] += aberto

    for cod, v in sorted(orfao.items(), key=lambda x: -abs(x[1])):
        controles.append(ctrl("INCONSISTÊNCIA", "Faturamento", "faturamento órfão (código vazio ou fora "
                              "do cadastro Projetos) — some do DRE sem aviso", cod, v))
    if canceladas:
        controles.append(ctrl("CONFIRMADO", "Faturamento", "NFs canceladas excluídas das somas",
                              str(canceladas), None))
    if apos_corte:
        controles.append(ctrl("CONFIRMADO", "Faturamento", "NFs emitidas após a data-base excluídas",
                              str(apos_corte), None))
    if texto_num:
        controles.append(ctrl("INCONSISTÊNCIA", "Faturamento", "valores gravados como texto (convertidos "
                              "pelo script — corrigir na origem)", str(texto_num), None))
    return {"colunas": {k: (header[v] if v is not None and v < len(header) else None) for k, v in
                        {"projeto": i_proj, "valor_bruto": i_bruto, "valor_recebido": i_receb,
                         "valor_liquido": i_liq, "vencimento": i_venc, "emissao": i_emis,
                         "cliente": i_cli, "situacao": i_sit}.items()},
            "colunas_retencao": [header[i] for i in i_ret],
            "linhas": len(corpo), "premissas": premissas}


# ------------------------------------------------------------------------ Cockpit

NATUREZAS = {
    "mo": ("salario", "mao de obra", "folha"),
    "materiais": ("materia",),
    "mobilizacao": ("mobiliza",),
}


def natureza_de(texto):
    n = norm(texto)
    for k, termos in NATUREZAS.items():
        if any(t in n for t in termos):
            return k
    return None


def aplicar_cockpit(pr, nat, valor, controles, cod):
    # Convencao AW: Salarios e Materiais vem positivos -> inverter; Mobilizacao ja vem negativa.
    if nat in ("mo", "materiais"):
        if valor < 0:
            controles.append(ctrl("INCONSISTÊNCIA", "Cockpit", f"{nat} já negativo no Cockpit (convenção "
                                  "é positivo) — mantido negativo, não invertido", cod, valor))
            v = valor
        else:
            v = -valor
        pr["cockpit_mo" if nat == "mo" else "cockpit_materiais"] += v
    elif nat == "mobilizacao":
        if valor > 0:
            controles.append(ctrl("INCONSISTÊNCIA", "Cockpit", "mobilização positiva (convenção é "
                                  "negativa) — mantida como veio", cod, valor))
        pr["cockpit_mobilizacao"] += valor
    else:
        pr["cockpit_outros"] += valor
    pr["no_cockpit"] = True


def ler_cockpit(ws, projetos, controles):
    header, corpo, _ = ler_tabela(ws)
    i_proj = col_index(header, "projeto", "obra", "centro de custo")
    if i_proj is None:
        controles.append(ctrl("PENDÊNCIA", "Cockpit", "layout do Cockpit não reconhecido (sem coluna "
                              "Projeto/Obra) — custo por obra não extraído", ws.title, None))
        return {"layout": "não reconhecido"}
    cols_nat = {i: natureza_de(h) for i, h in enumerate(header) if natureza_de(h)}
    i_valor = col_index(header, "valor", "total", "realizado")
    i_desc = None
    if not cols_nat and i_valor is not None:
        for i in range(len(header)):
            if i in (i_proj, i_valor):
                continue
            if any(natureza_de(get(r, i)) for r in corpo[:200]):
                i_desc = i
                break
    layout = "largo" if cols_nat else ("longo" if i_desc is not None else None)
    if layout is None:
        controles.append(ctrl("PENDÊNCIA", "Cockpit", "não achei colunas Salários/Materiais/Mobilização "
                              "nem coluna de natureza + valor — confirmar layout", ws.title, None))
        return {"layout": "não reconhecido"}
    for r in corpo:
        p = parse_projeto(get(r, i_proj))
        if not p or not p["codigo"]:
            continue
        pr = projetos.setdefault(p["codigo"], novo_projeto())
        if layout == "largo":
            for i, nat in cols_nat.items():
                v, _ = parse_numero(get(r, i))
                if v:
                    aplicar_cockpit(pr, nat, v, controles, p["codigo"])
        else:
            v, _ = parse_numero(get(r, i_valor))
            if v:
                aplicar_cockpit(pr, natureza_de(get(r, i_desc)), v, controles, p["codigo"])
    return {"layout": layout,
            "colunas_natureza": {header[i]: n for i, n in cols_nat.items()} if cols_nat else None,
            "coluna_natureza": header[i_desc] if i_desc is not None else None}


# ------------------------------------------------------------------------ Planilhao

def ler_planilhao(ws, projetos, controles):
    header, corpo, _ = ler_tabela(ws)
    i_un = col_index(header, "unidade de negocio", "unidade")
    i_pai = col_index(header, "plano de contas pai")
    i_conta = col_index(header, "conta corrente", "conta")
    i_valor = col_index(header, "valor")
    i_proj = col_index(header, "projeto")
    i_mes = col_index(header, "mes")
    i_ano = col_index(header, "ano")
    i_data = col_index(header, "data incorrido", "data")
    if None in (i_un, i_conta, i_valor):
        controles.append(ctrl("PENDÊNCIA", "Planilhão", "sem Unidade/Conta Corrente/Valor — caixa real x "
                              "virtual não calculado", ws.title, None))
        return None
    caixa = defaultdict(lambda: {"real": 0.0, "virtual": 0.0, "situacao": None})
    virtual_mes = defaultdict(float)
    virtual_conta = defaultdict(float)
    for r in corpo:
        v, _ = parse_numero(get(r, i_valor))
        if v is None:
            continue
        un, sit = classificar_unidade(get(r, i_un))
        conta = get(r, i_conta)
        c = caixa[un]
        c["situacao"] = sit
        if eh_conta_bancaria(conta):
            c["real"] += v
        else:
            c["virtual"] += v
            ano, mes = get(r, i_ano), get(r, i_mes)
            if (ano is None or mes is None) and i_data is not None:
                d = parse_data(get(r, i_data))
                if d:
                    ano, mes = d.year, d.month
            virtual_mes[(ano, mes)] += v
            virtual_conta[str(conta).strip() if conta is not None else "(vazia)"] += v
        if i_proj is not None:
            p = parse_projeto(get(r, i_proj))
            if p and p["codigo"]:
                pr = projetos.setdefault(p["codigo"], novo_projeto())
                pr["no_planilhao"] = True
                pai = norm(get(r, i_pai)) if i_pai is not None else ""
                if pai.startswith("001"):
                    pr["planilhao_receita"] += v
                elif pai.startswith("003") or pai.startswith("005"):
                    pr["planilhao_custo"] += v
    return {"caixa_por_unidade": caixa, "virtual_por_mes": virtual_mes,
            "virtual_por_conta": virtual_conta}


# ------------------------------------------------------------------------ utilitarios

def ctrl(rotulo, fonte, descricao, chave, valor):
    return {"rotulo": rotulo, "fonte": fonte, "descricao": descricao, "chave": chave,
            "valor": round(valor, 2) if isinstance(valor, float) else valor}


def fechar_projetos(projetos, controles, tol):
    for cod, pr in projetos.items():
        pr["a_receber"] = pr["faturado"] - pr["recebido"] - pr["retencoes"]
        pr["custo_cockpit"] = (pr["cockpit_mo"] + pr["cockpit_materiais"] + pr["cockpit_mobilizacao"]
                               + pr["cockpit_outros"])
        custo = pr["custo_cockpit"] if pr["no_cockpit"] else (pr["planilhao_custo"] if pr["no_planilhao"] else None)
        pr["custo_incorrido"] = custo
        pr["fonte_custo"] = "Cockpit" if pr["no_cockpit"] else ("Planilhão" if pr["no_planilhao"] else None)
        if custo is not None and pr["faturado"]:
            pr["resultado"] = pr["faturado"] + custo
            pr["margem_real"] = pr["resultado"] / pr["faturado"]
        else:
            pr["resultado"] = (pr["faturado"] + custo) if custo is not None else None
            pr["margem_real"] = None
        alertas = []
        if pr["recebido"] - pr["faturado"] > tol:
            alertas.append("recebido > faturado")
            controles.append(ctrl("INCONSISTÊNCIA", "Faturamento", "recebido maior que faturado "
                                  "(adiantamento? NF faltando? vínculo errado?)", cod,
                                  pr["recebido"] - pr["faturado"]))
        if pr["resultado"] is not None and pr["resultado"] < -tol:
            alertas.append("resultado negativo")
            controles.append(ctrl("ACHADO CRÍTICO", "Projeto", "resultado realizado negativo "
                                  "(faturado + custo incorrido)", cod, pr["resultado"]))
        if pr["custo_incorrido"] is not None and pr["faturado"] == 0 and pr["recebido"] == 0:
            alertas.append("custo sem faturamento")
        if pr["no_cockpit"] and pr["no_planilhao"]:
            dif = pr["custo_cockpit"] - pr["planilhao_custo"]
            pr["dif_cockpit_planilhao"] = dif
            if abs(dif) > tol:
                alertas.append("Cockpit ≠ Planilhão")
                controles.append(ctrl("INCONSISTÊNCIA", "Cockpit x Planilhão", "custo por obra diverge "
                                      "(rateio, sinal ou competência)", cod, dif))
        if not pr["retencao_conhecida"] and pr["a_receber"] > tol and pr["recebido"] > 0:
            alertas.append("CR inclui retenções (PREMISSA)")
        if pr["aging"].get(">90", 0) > tol:
            alertas.append("CR > 90 dias")
        pr["alertas"] = alertas


def somar_unidades(projetos, caixa):
    un = defaultdict(lambda: defaultdict(float))
    sit = {}
    for cod, pr in projetos.items():
        u = pr["unidade"] or "SEM UNIDADE"
        sit[u] = pr["situacao_unidade"] or SITUACAO_FORA
        d = un[u]
        d["projetos"] += 1
        for k in ("faturado", "recebido", "retencoes", "a_receber"):
            d[k] += pr[k]
        if pr["custo_incorrido"] is not None:
            d["custo_incorrido"] += pr["custo_incorrido"]
        for f, v in pr["aging"].items():
            d["aging " + f] += v
    if caixa:
        for u, c in caixa["caixa_por_unidade"].items():
            sit.setdefault(u, c["situacao"])
            un[u]["caixa_real"] += c["real"]
            un[u]["caixa_virtual"] += c["virtual"]
    for u, d in un.items():
        d["resultado"] = d["faturado"] + d["custo_incorrido"]
        d["margem_real"] = (d["resultado"] / d["faturado"]) if d["faturado"] else None
        d["caixa_gerencial"] = d["caixa_real"] + d["caixa_virtual"]
    return un, sit


def consolidar(un, sit):
    grupos = defaultdict(lambda: defaultdict(float))
    for u, d in un.items():
        g = sit.get(u) or SITUACAO_FORA
        for k, v in d.items():
            if isinstance(v, (int, float)) and v is not None and k != "margem_real":
                grupos[g][k] += v
    for g in grupos.values():
        g["margem_real"] = (g["resultado"] / g["faturado"]) if g.get("faturado") else None
    return grupos


# ------------------------------------------------------------------------ saida

ORDEM_SIT = [SITUACAO_ATIVA, SITUACAO_DESATIVADA, SITUACAO_A_PARTE, SITUACAO_FORA]


def imprimir_resumo(projetos, un, sit, grupos, controles, caixa, meta):
    print("=" * 78)
    print("PAINEL AW — resumo (detalhe no JSON/XLSX)")
    print("=" * 78)
    for p in meta.get("premissas", []):
        print(f"[PREMISSA] {p}")
    criticos = [c for c in controles if c["rotulo"] == "ACHADO CRÍTICO"]
    if criticos:
        print(f"\n>> ACHADOS CRÍTICOS ({len(criticos)})")
        for c in sorted(criticos, key=lambda c: c["valor"] or 0)[:10]:
            print(f"   {c['chave']}: {c['descricao']} = {fmt_brl(c['valor'])}")
    print("\n>> CONSOLIDADO POR SITUAÇÃO DA UNIDADE")
    cab = f"{'Grupo':<20}{'Faturado':>18}{'Recebido':>18}{'A receber':>18}{'Custo inc.':>18}"
    print(cab)
    for g in ORDEM_SIT:
        if g in grupos:
            d = grupos[g]
            print(f"{g:<20}{fmt_brl(d['faturado']):>18}{fmt_brl(d['recebido']):>18}"
                  f"{fmt_brl(d['a_receber']):>18}{fmt_brl(d['custo_incorrido']):>18}")
    print("\n>> POR UNIDADE (ativas primeiro)")
    print(f"{'Unidade':<26}{'Faturado':>16}{'A receber':>16}{'Resultado':>16}{'Caixa real':>16}{'Caixa virt.':>16}")
    for u in sorted(un, key=lambda x: (ORDEM_SIT.index(sit.get(x) or SITUACAO_FORA), x)):
        d = un[u]
        print(f"{u[:25]:<26}{fmt_brl(d['faturado']):>16}{fmt_brl(d['a_receber']):>16}"
              f"{fmt_brl(d['resultado']):>16}{fmt_brl(d['caixa_real']):>16}{fmt_brl(d['caixa_virtual']):>16}")
    if caixa:
        print("\n>> CONTAS VIRTUAIS NO CONSOLIDADO (devem somar ≈ 0 por mês)")
        for (a, m), v in sorted(caixa["virtual_por_mes"].items(), key=lambda x: (str(x[0][0]), str(x[0][1]))):
            flag = "OK" if abs(v) <= meta["tolerancia"] else "NÃO ZERA"
            print(f"   {m}/{a}: {fmt_brl(v)}  {flag}")
    print(f"\n>> CONTROLES: {len(controles)} (INCONSISTÊNCIA: "
          f"{sum(c['rotulo'] == 'INCONSISTÊNCIA' for c in controles)}, PENDÊNCIA: "
          f"{sum(c['rotulo'] == 'PENDÊNCIA' for c in controles)})")
    for c in controles[:25]:
        val = f" = {fmt_brl(c['valor'])}" if isinstance(c["valor"], float) else ""
        print(f"   [{c['rotulo']}] {c['fonte']}: {c['descricao']} — {c['chave']}{val}")
    if len(controles) > 25:
        print(f"   ... +{len(controles) - 25} no JSON/XLSX")


def gravar_xlsx(caminho, projetos, un, sit, grupos, controles, caixa, meta):
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    brl = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
    pct = "0.0%"
    negrito = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="0F3D3D")

    def aba(nome, cab, linhas, fmt):
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = negrito, fundo
        for l in linhas:
            ws.append(l)
        for j, f in enumerate(fmt, start=1):
            if f:
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    row[0].number_format = f
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                60, max(12, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    leia = wb.active
    leia.title = "LEIA-ME"
    for l in [["Painel AW — gerado por controller-empreiteira/scripts/painel_aw.py"],
              ["Arquivo de origem", meta["arquivo"]], ["Data-base", meta.get("data_base") or "não informada"],
              ["Convenção", "Receita positiva, custo negativo. Cockpit: Salários/Materiais invertidos; "
                            "Mobilização como veio."],
              ["Caixa real", "movimentos do Planilhão em conta bancária (NNNNN-N)"],
              ["Caixa virtual", "movimentos em conta gerencial (não é caixa)"],
              ["Rótulos", "CONFIRMADO · PREMISSA · PENDÊNCIA · INCONSISTÊNCIA · ACHADO CRÍTICO"]] + \
             [["PREMISSA", p] for p in meta.get("premissas", [])]:
        leia.append(l)
    leia.column_dimensions["A"].width = 22
    leia.column_dimensions["B"].width = 110

    cab = ["Código", "Descritivo", "Unidade", "Situação unidade", "Status", "No cadastro", "NFs",
           "Faturado", "Recebido", "Retenções", "A receber", "Retenção conhecida"] + \
          [f"Aging {f}" for f in FAIXAS_AGING] + \
          ["Cockpit MO", "Cockpit Materiais", "Cockpit Mobilização", "Cockpit Outros", "Custo Cockpit",
           "Custo Planilhão", "Dif Cockpit−Planilhão", "Fonte custo", "Custo incorrido", "Resultado",
           "Margem real", "Alertas"]
    fmt = [None] * 7 + [brl] * 4 + [None] + [brl] * len(FAIXAS_AGING) + [brl] * 7 + [None, brl, brl, pct, None]
    linhas = []
    for cod in sorted(projetos, key=lambda c: (projetos[c]["unidade"] or "~", c)):
        p = projetos[cod]
        linhas.append([cod, p["descritivo"], p["unidade"], p["situacao_unidade"], p["status"],
                       "sim" if p["no_cadastro"] else "NÃO", p["n_nf"], p["faturado"], p["recebido"],
                       p["retencoes"], p["a_receber"], "sim" if p["retencao_conhecida"] else "não (PREMISSA)"] +
                      [p["aging"].get(f, 0.0) for f in FAIXAS_AGING] +
                      [p["cockpit_mo"], p["cockpit_materiais"], p["cockpit_mobilizacao"], p["cockpit_outros"],
                       p["custo_cockpit"], p["planilhao_custo"], p.get("dif_cockpit_planilhao"),
                       p["fonte_custo"], p["custo_incorrido"], p["resultado"], p["margem_real"],
                       "; ".join(p["alertas"])])
    ws = aba("PROJETOS", cab, linhas, fmt)
    ws.auto_filter.ref = ws.dimensions

    cab_u = ["Unidade", "Situação", "Projetos", "Faturado", "Recebido", "Retenções", "A receber",
             "Custo incorrido", "Resultado", "Margem real", "Caixa real", "Caixa virtual", "Caixa gerencial"] + \
            [f"Aging {f}" for f in FAIXAS_AGING]
    fmt_u = [None, None, None] + [brl] * 6 + [pct] + [brl] * 3 + [brl] * len(FAIXAS_AGING)
    lin_u = []
    for u in sorted(un, key=lambda x: (ORDEM_SIT.index(sit.get(x) or SITUACAO_FORA), x)):
        d = un[u]
        lin_u.append([u, sit.get(u), int(d["projetos"]), d["faturado"], d["recebido"], d["retencoes"],
                      d["a_receber"], d["custo_incorrido"], d["resultado"], d["margem_real"], d["caixa_real"],
                      d["caixa_virtual"], d["caixa_gerencial"]] + [d.get("aging " + f, 0.0) for f in FAIXAS_AGING])
    for g in ORDEM_SIT:
        if g in grupos:
            d = grupos[g]
            lin_u.append([f"TOTAL {g.upper()}", g, int(d.get("projetos", 0)), d["faturado"], d["recebido"],
                          d["retencoes"], d["a_receber"], d["custo_incorrido"], d["resultado"], d["margem_real"],
                          d["caixa_real"], d["caixa_virtual"], d["caixa_gerencial"]] +
                         [d.get("aging " + f, 0.0) for f in FAIXAS_AGING])
    aba("UNIDADES", cab_u, lin_u, fmt_u)

    if caixa:
        lin_v = [[f"{m}/{a}", v, "OK" if abs(v) <= meta["tolerancia"] else "NÃO ZERA"]
                 for (a, m), v in sorted(caixa["virtual_por_mes"].items(), key=lambda x: (str(x[0][0]), str(x[0][1])))]
        aba("VIRTUAL_MES", ["Mês/Ano", "Soma contas virtuais (consolidado)", "Teste zero"], lin_v, [None, brl, None])
        lin_c = [[k, v] for k, v in sorted(caixa["virtual_por_conta"].items(), key=lambda x: -abs(x[1]))]
        aba("VIRTUAL_CONTA", ["Conta virtual", "Saldo no período (consolidado)"], lin_c, [None, brl])

    aba("CONTROLES", ["Rótulo", "Fonte", "Descrição", "Chave", "Valor"],
        [[c["rotulo"], c["fonte"], c["descricao"], str(c["chave"]), c["valor"] if isinstance(c["valor"], float) else None]
         for c in controles], [None, None, None, None, brl])
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo")
    ap.add_argument("--planilhao", help="arquivo separado com o Planilhão (se não estiver no mesmo)")
    ap.add_argument("--data-base", help="dd/mm/aaaa — corte de emissão e referência do aging")
    ap.add_argument("--aba-faturamento")
    ap.add_argument("--aba-projetos")
    ap.add_argument("--aba-cockpit")
    ap.add_argument("--aba-planilhao")
    ap.add_argument("--tolerancia", type=float, default=1.0)
    ap.add_argument("--json")
    ap.add_argument("--xlsx")
    args = ap.parse_args()

    data_base = parse_data(args.data_base) if args.data_base else None
    if args.data_base and not data_base:
        sys.exit(f"Data-base inválida: {args.data_base} (use dd/mm/aaaa)")

    wb = openpyxl.load_workbook(args.arquivo, read_only=True, data_only=True)
    projetos, controles, meta = {}, [], {"arquivo": args.arquivo, "tolerancia": args.tolerancia,
                                         "data_base": args.data_base, "premissas": [], "abas": {}}
    if not data_base:
        meta["premissas"].append("sem --data-base: aging por vencimento fica 'Sem vencimento' e não há "
                                 "corte de emissão.")

    ws_p = wb[args.aba_projetos] if args.aba_projetos else achar_aba(wb, "projetos", "projeto")
    ws_f = wb[args.aba_faturamento] if args.aba_faturamento else achar_aba(wb, "faturamento")
    ws_c = wb[args.aba_cockpit] if args.aba_cockpit else achar_aba(wb, "cockpit")
    if args.planilhao:
        wb2 = openpyxl.load_workbook(args.planilhao, read_only=True, data_only=True)
        ws_pl = wb2[args.aba_planilhao] if args.aba_planilhao else (achar_aba(wb2, "planilh", "incorrid") or wb2.worksheets[0])
    else:
        ws_pl = wb[args.aba_planilhao] if args.aba_planilhao else achar_aba(wb, "planilh", "incorrid")

    if ws_p is not None:
        meta["abas"]["projetos"] = {"aba": ws_p.title, **ler_projetos(ws_p, projetos, controles)}
    else:
        controles.append(ctrl("PENDÊNCIA", "Projetos", "aba Projetos não encontrada — todo faturamento "
                              "fica sem vínculo de unidade", args.arquivo, None))
    if ws_f is not None:
        info = ler_faturamento(ws_f, projetos, controles, data_base)
        meta["premissas"] += info.pop("premissas")
        meta["abas"]["faturamento"] = {"aba": ws_f.title, **info}
    else:
        controles.append(ctrl("PENDÊNCIA", "Faturamento", "aba Faturamento não encontrada", args.arquivo, None))
    if ws_c is not None:
        meta["abas"]["cockpit"] = {"aba": ws_c.title, **ler_cockpit(ws_c, projetos, controles)}
    caixa = None
    if ws_pl is not None:
        caixa = ler_planilhao(ws_pl, projetos, controles)
        meta["abas"]["planilhao"] = ws_pl.title
        if caixa:
            for (a, m), v in caixa["virtual_por_mes"].items():
                if abs(v) > args.tolerancia:
                    controles.append(ctrl("INCONSISTÊNCIA", "Planilhão", "contas virtuais não zeram no "
                                          "consolidado (perna faltante ou em outro mês)", f"{m}/{a}", v))
    else:
        meta["premissas"].append("sem Planilhão: caixa real x virtual e amarração Cockpit x Planilhão "
                                 "não calculados.")

    fechar_projetos(projetos, controles, args.tolerancia)
    un, sit = somar_unidades(projetos, caixa)
    grupos = consolidar(un, sit)
    ordem = {"ACHADO CRÍTICO": 0, "INCONSISTÊNCIA": 1, "PENDÊNCIA": 2, "CONFIRMADO": 3}
    controles.sort(key=lambda c: (ordem.get(c["rotulo"], 9), -abs(c["valor"]) if isinstance(c["valor"], float) else 0))

    imprimir_resumo(projetos, un, sit, grupos, controles, caixa, meta)

    if args.json:
        def limpo(pr):
            d = dict(pr)
            d["aging"] = dict(pr["aging"])
            d["clientes"] = dict(pr["clientes"])
            return d
        saida = {"meta": meta, "consolidado": {g: dict(v) for g, v in grupos.items()},
                 "unidades": {u: {**dict(d), "situacao": sit.get(u)} for u, d in un.items()},
                 "projetos": {c: limpo(p) for c, p in projetos.items()}, "controles": controles}
        if caixa:
            saida["virtual_por_mes"] = {f"{m}/{a}": v for (a, m), v in caixa["virtual_por_mes"].items()}
            saida["virtual_por_conta"] = dict(caixa["virtual_por_conta"])
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(saida, f, ensure_ascii=False, indent=2, default=str)
        print(f"\nJSON salvo em {args.json}", file=sys.stderr)
    if args.xlsx:
        gravar_xlsx(args.xlsx, projetos, un, sit, grupos, controles, caixa, meta)
        print(f"Planilha de trabalho salva em {args.xlsx} (arquivo de origem não foi alterado)", file=sys.stderr)


if __name__ == "__main__":
    main()
